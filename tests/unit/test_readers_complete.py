"""Every reader gives the whole document: nothing a person would read is left out, and every part says where it is.

A Word document is read in order, tables as rows where they sit, lists with
their markers, footnotes after the paragraph that calls them, text boxes
once; its pages are where Word marked them, and a section that numbers its
own pages gives the printed numbers. A PDF's bookmarks are its headings. A
deck's speaker notes follow their slide. A web page's tables are rows. And
every file's own title, where it has one, is on every chunk.
"""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

import pytest

from vectrixdb.ingest import _useful_title, load, load_bytes, prepare_document

NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
)


def docx_module():
    return pytest.importorskip("docx")


def raw(fragment):
    """A paragraph written as WordprocessingML."""
    from docx.oxml import parse_xml

    return parse_xml(f"<w:p {NS}>{fragment}</w:p>")


def put(document, element):
    """Put an element at the end of the body, before its section properties."""
    body = document.element.body
    if body.sectPr is not None:
        body.sectPr.addprevious(element)
    else:
        body.append(element)


def new_page_at_start(paragraph):
    """Word's mark that a page began with this paragraph."""
    from docx.oxml import parse_xml

    run = parse_xml(f"<w:r {NS}><w:lastRenderedPageBreak/></w:r>")
    properties = paragraph._p.pPr
    if properties is not None:
        properties.addnext(run)
    else:
        paragraph._p.insert(0, run)


def saved(document, tmp_path, name="policy.docx"):
    path = tmp_path / name
    document.save(str(path))
    return path


# ================================================================ Word ===


class TestAWordDocumentIsReadWhole:
    def test_a_table_is_read_where_it_sits_as_rows(self, tmp_path):
        docx = docx_module()
        d = docx.Document()
        d.add_heading("Terms", level=1)
        d.add_paragraph("Payment is due within thirty days.")
        t = d.add_table(rows=3, cols=3)
        for r, row in enumerate(
            [["Region", "Revenue", "Growth"], ["EMEA", "1200", "4%"], ["APAC", "950", "0%"]]
        ):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
        d.add_paragraph("Figures are unaudited.")
        text = load(saved(d, tmp_path)).text
        assert text == (
            "# Terms\n\nPayment is due within thirty days.\n\n"
            "Region: EMEA; Revenue: 1200; Growth: 4%\nRegion: APAC; Revenue: 950; Growth: 0%\n\n"
            "Figures are unaudited."
        )

    def test_a_merged_cell_says_its_value_in_every_row(self, tmp_path):
        docx = docx_module()
        d = docx.Document()
        t = d.add_table(rows=3, cols=3)
        for r, row in enumerate(
            [["Region", "Quarter", "Revenue"], ["EMEA", "Q1", "100"], ["", "Q2", "120"]]
        ):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
        t.cell(1, 0).merge(t.cell(2, 0))
        lines = load(saved(d, tmp_path)).text.splitlines()
        assert lines == [
            "Region: EMEA; Quarter: Q1; Revenue: 100",
            "Region: EMEA; Quarter: Q2; Revenue: 120",
        ]

    def test_a_cell_spanning_two_columns_keeps_the_rest_in_line(self, tmp_path):
        docx = docx_module()
        d = docx.Document()
        t = d.add_table(rows=2, cols=3)
        for c, value in enumerate(["Region", "Plan", "Actual"]):
            t.cell(0, c).text = value
        t.cell(1, 0).text = "EMEA"
        t.cell(1, 2).text = "130"
        t.cell(1, 0).merge(t.cell(1, 1))
        assert load(saved(d, tmp_path)).text == "Region: EMEA; Actual: 130"

    def test_a_table_inside_a_cell_is_read_as_its_words(self, tmp_path):
        docx = docx_module()
        d = docx.Document()
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Mode", "Speed"
        t.cell(1, 0).text = "hybrid"
        inner = t.cell(1, 1).add_table(rows=1, cols=2)
        inner.cell(0, 0).text, inner.cell(0, 1).text = "fast", "on a laptop"
        assert load(saved(d, tmp_path)).text == "Mode: hybrid; Speed: fast on a laptop"

    def test_list_items_keep_their_bullets_and_numbers(self, tmp_path):
        docx = docx_module()
        d = docx.Document()
        d.add_paragraph("A reminder after ten days.", style="List Bullet")
        d.add_paragraph("A second after twenty.", style="List Bullet")
        for step in ("Open the account.", "Sign.", "Fund it."):
            d.add_paragraph(step, style="List Number")
        assert load(saved(d, tmp_path)).text.split("\n\n") == [
            "- A reminder after ten days.",
            "- A second after twenty.",
            "1. Open the account.",
            "2. Sign.",
            "3. Fund it.",
        ]

    def test_footnotes_and_endnotes_follow_the_paragraph_that_calls_them(self, tmp_path):
        docx = docx_module()
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.opc.packuri import PackURI
        from docx.opc.part import Part

        d = docx.Document()
        put(
            d,
            raw(
                '<w:r><w:t xml:space="preserve">Revenue grew</w:t></w:r><w:r><w:footnoteReference w:id="7"/></w:r><w:r><w:t>.</w:t></w:r>'
            ),
        )
        put(
            d,
            raw(
                '<w:r><w:t xml:space="preserve">Costs were flat</w:t></w:r><w:r><w:endnoteReference w:id="2"/></w:r>'
            ),
        )
        for kind, rel, content_type, note in (
            (
                "footnote",
                RT.FOOTNOTES,
                "footnotes",
                '<w:footnote w:id="7"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> Unaudited.</w:t></w:r></w:p></w:footnote>',
            ),
            (
                "endnote",
                RT.ENDNOTES,
                "endnotes",
                '<w:endnote w:id="2"><w:p><w:r><w:t>See note 5.</w:t></w:r></w:p></w:endnote>',
            ),
        ):
            separator = f'<w:{kind} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{kind}>'
            blob = f"<w:{kind}s {NS}>{separator}{note}</w:{kind}s>".encode()
            part = Part(
                PackURI(f"/word/{content_type}.xml"),
                f"application/vnd.openxmlformats-officedocument.wordprocessingml.{content_type}+xml",
                blob,
                d.part.package,
            )
            d.part.relate_to(part, rel)
        assert load(saved(d, tmp_path)).text.split("\n\n") == [
            "Revenue grew[^1].",
            "[^1]: Unaudited.",
            "Costs were flat[^e1]",
            "[^e1]: See note 5.",
        ]

    def test_a_text_box_is_read_once_after_its_paragraph(self, tmp_path):
        docx = docx_module()
        d = docx.Document()
        box = "<w:txbxContent><w:p><w:r><w:t>Sidebar: rates rose.</w:t></w:r></w:p></w:txbxContent>"
        put(
            d,
            raw(
                "<w:r><w:t>Anchor paragraph.</w:t></w:r><w:r><mc:AlternateContent>"
                f'<mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic><a:graphicData><wps:wsp><wps:txbx>{box}</wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>'
                f"<mc:Fallback><w:pict><v:shape><v:textbox>{box}</v:textbox></v:shape></w:pict></mc:Fallback>"
                "</mc:AlternateContent></w:r>"
            ),
        )
        assert load(saved(d, tmp_path)).text == "Anchor paragraph.\n\nSidebar: rates rose."

    def test_what_is_shown_is_read_and_what_is_not_is_not(self, tmp_path):
        """A content control's words, an insertion and a field's result are shown; a deletion and a field's code are not."""
        docx = docx_module()
        from docx.oxml import parse_xml

        d = docx.Document()
        put(
            d,
            parse_xml(
                f"<w:sdt {NS}><w:sdtPr/><w:sdtContent><w:p><w:r><w:t>Clause 4 applies.</w:t></w:r></w:p></w:sdtContent></w:sdt>"
            ),
        )
        put(
            d,
            raw(
                '<w:r><w:t xml:space="preserve">Rate: </w:t></w:r>'
                "<w:del><w:r><w:delText>5%</w:delText><w:tab/></w:r></w:del><w:ins><w:r><w:t>6%</w:t></w:r></w:ins>"
                '<w:r><w:t xml:space="preserve">, page </w:t></w:r>'
                '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> PAGE </w:instrText></w:r>'
                '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>3</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r>'
            ),
        )
        assert load(saved(d, tmp_path)).text == "Clause 4 applies.\n\nRate: 6%, page 3"

    def test_the_title_is_the_files_own_else_its_title_paragraph(self, tmp_path):
        docx = docx_module()
        d = docx.Document()
        d.add_paragraph("Lending Policy 2025", style="Title")
        d.add_heading("Terms", level=1)
        doc = load(saved(d, tmp_path))
        assert doc.metadata["title"] == "Lending Policy 2025"
        assert doc.text.startswith("# Lending Policy 2025\n\n# Terms"), (
            "and the first part is not left without a heading"
        )
        d.core_properties.title = "Lending policy, 2025 edition"
        assert (
            load(saved(d, tmp_path, "b.docx")).metadata["title"] == "Lending policy, 2025 edition"
        )


class TestAWordDocumentsPages:
    def test_pages_where_word_marked_them(self, tmp_path):
        docx = docx_module()
        d = docx.Document()
        d.add_paragraph("Page one text.")
        b = d.add_paragraph("Page two starts here.")
        new_page_at_start(b)
        put(
            d,
            raw(
                '<w:r><w:t xml:space="preserve">Still page two. </w:t></w:r><w:r><w:lastRenderedPageBreak/><w:t>Page three begins mid-paragraph.</w:t></w:r>'
            ),
        )
        put(d, raw('<w:r><w:t>Before a manual break.</w:t><w:br w:type="page"/></w:r>'))
        f = d.add_paragraph("After the manual break.")
        new_page_at_start(f)
        doc = load(saved(d, tmp_path))
        text = doc.text
        assert doc.pages == [
            (0, 1),
            (text.index("Page two starts"), 2),
            (text.index("Page three begins"), 3),
            (text.index("After the manual"), 4),
        ], "the manual break and the mark Word wrote after it are one new page"
        assert doc.metadata["pages"] == 4

    def test_a_page_that_begins_inside_a_table_begins_at_its_row(self, tmp_path):
        docx = docx_module()
        from docx.oxml import parse_xml

        d = docx.Document()
        d.add_paragraph("Before the table.")
        t = d.add_table(rows=3, cols=2)
        for r, row in enumerate([["Region", "Revenue"], ["EMEA", "1200"], ["APAC", "950"]]):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
        t.cell(2, 0).paragraphs[0]._p.insert(
            0, parse_xml(f"<w:r {NS}><w:lastRenderedPageBreak/></w:r>")
        )
        doc = load(saved(d, tmp_path))
        assert doc.pages == [(0, 1), (doc.text.index("Region: APAC"), 2)]

    def test_a_manual_break_and_words_own_mark_a_space_apart_are_one_page(self, tmp_path):
        docx = docx_module()
        d = docx.Document()
        put(
            d,
            raw(
                '<w:r><w:t>Before.</w:t><w:br w:type="page"/><w:t xml:space="preserve"> </w:t><w:lastRenderedPageBreak/><w:t>After.</w:t></w:r>'
            ),
        )
        assert len(load(saved(d, tmp_path)).pages) == 2

    def test_a_file_word_never_laid_out_has_no_pages(self, tmp_path):
        """Pages counted from manual breaks alone would be wrong wherever text ran onto a new page by itself."""
        docx = docx_module()
        d = docx.Document()
        d.add_paragraph("One.")
        d.add_page_break()
        d.add_paragraph("Two.")
        doc = load(saved(d, tmp_path))
        assert doc.pages == [] and "pages" not in doc.metadata

    def test_a_section_that_numbers_its_own_pages(self, tmp_path):
        docx = docx_module()
        from docx.enum.section import WD_SECTION
        from docx.oxml import parse_xml

        d = docx.Document()
        d.add_paragraph("Preface, first page.")
        new_page_at_start(d.add_paragraph("Preface, second page."))
        d.add_section(WD_SECTION.NEW_PAGE)
        new_page_at_start(d.add_paragraph("Chapter one."))
        new_page_at_start(d.add_paragraph("Chapter one, next page."))
        d.sections[0]._sectPr.append(parse_xml(f'<w:pgNumType {NS} w:fmt="lowerRoman"/>'))
        d.sections[1]._sectPr.append(parse_xml(f'<w:pgNumType {NS} w:start="1"/>'))
        doc = load(saved(d, tmp_path))
        assert [n for _o, n in doc.pages] == [1, 2, 3, 4]
        assert doc.page_labels == {1: "i", 2: "ii", 3: "1", 4: "2"}

    def test_a_word_files_chunks_carry_their_pages(self, tmp_path):
        docx = docx_module()
        d = docx.Document()
        d.add_heading("Terms", level=1)
        d.add_paragraph("Payment is due within thirty days of the invoice date, every time.")
        new_page_at_start(d.add_heading("Late fees", level=1))
        d.add_paragraph("Interest accrues monthly on any overdue balance, and a reminder follows.")
        prepared = prepare_document(
            load(saved(d, tmp_path)), "policy.docx", chunk="markdown", chunk_size=1000, overlap=0
        )
        late = next(m for m in prepared.metadata if m.get("heading") == "Late fees")
        assert (late["page"], late["_vx_citation"], late["_vx_readable_citation"]) == (
            2,
            "policy.docx#page=2",
            "policy.docx, p. 2",
        )


# ================================================================ PDF ===


def text_pdf(pages, outline=(), title=None):
    """A PDF with these lines on its pages, these bookmarks, ``(title, page index, parent title)``, and this title."""
    pypdf = pytest.importorskip("pypdf")
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = pypdf.PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    for lines in pages:
        page = writer.add_blank_page(width=612, height=792)
        escaped = [
            line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)") for line in lines
        ]
        stream = DecodedStreamObject()
        stream.set_data(
            "\n".join(
                ["BT", "/F1 12 Tf", "16 TL", "72 720 Td"]
                + [f"({line}) Tj T*" for line in escaped]
                + ["ET"]
            ).encode("latin-1")
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        page.replace_contents(stream)
    made = {}
    for name, index, parent in outline:
        made[name] = writer.add_outline_item(name, index, parent=made.get(parent))
    if title:
        writer.add_metadata({"/Title": title})
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


REPORT = [
    ["Annual report 2025", "Who we are and what we do, in brief."],
    [
        "Our results were strong this year, and every region grew.",
        "Message from the CEO",
        "Dear shareholders, a good year for all of us.",
    ],
    ["Our strategy", "Three pillars guide the next five years of work."],
]
OUTLINE = [
    ("Annual report 2025", 0, None),
    ("Message from the CEO", 1, "Annual report 2025"),
    ("Our strategy", 2, "Annual report 2025"),
]


class TestAPdfsBookmarksAreItsHeadings:
    def test_each_bookmark_is_a_heading_where_its_title_is_written(self, tmp_path):
        path = tmp_path / "report.pdf"
        path.write_bytes(text_pdf(REPORT, OUTLINE, title="Annual Report 2025"))
        doc = load(path)
        assert [(title, level) for _o, title, level in doc.headings] == [
            ("Annual report 2025", 1),
            ("Message from the CEO", 2),
            ("Our strategy", 2),
        ]
        message = next(o for o, t, _l in doc.headings if t == "Message from the CEO")
        assert doc.text[message:].startswith("Message from the CEO"), (
            "at its title's line, not at the top of its page"
        )
        assert doc.metadata["title"] == "Annual Report 2025"

    def test_the_chunks_carry_their_section_and_the_title(self, tmp_path):
        path = tmp_path / "report.pdf"
        path.write_bytes(text_pdf(REPORT, OUTLINE, title="Annual Report 2025"))
        prepared = prepare_document(
            load(path), "report.pdf", chunk="markdown", chunk_size=1000, overlap=0
        )
        section = {m["heading"]: m for m in prepared.metadata}
        assert set(section) == {"Annual report 2025", "Message from the CEO", "Our strategy"}, (
            "cut at the sections"
        )
        assert section["Message from the CEO"]["page"] == 2
        assert all(m["title"] == "Annual Report 2025" for m in prepared.metadata)

    def test_a_bookmark_whose_title_is_not_on_its_page_starts_at_the_page(self, tmp_path):
        path = tmp_path / "report.pdf"
        path.write_bytes(text_pdf(REPORT, [("Chapter two", 2, None)]))
        doc = load(path)
        start = next(o for o, n in doc.pages if n == 3)
        assert doc.headings == [(start, "Chapter two", 1)]

    def test_a_pdf_without_bookmarks_or_a_title_is_as_before(self, tmp_path):
        path = tmp_path / "report.pdf"
        path.write_bytes(text_pdf(REPORT))
        doc = load(path)
        assert doc.headings == [] and "title" not in doc.metadata

    def test_read_in_pieces_the_bookmarks_and_title_come_from_the_original(self):
        """A piece cut from a PDF has no bookmarks and no title, so they are read from the whole file."""
        from vectrixdb.extract import batched

        def service(data, name, source=None):
            import pypdf

            pages = [page.extract_text() or "" for page in pypdf.PdfReader(io.BytesIO(data)).pages]
            offsets, offset = [], 0
            for n, text in enumerate(pages, start=1):
                offsets.append([offset, n])
                offset += len(text) + 2
            return {"text": "\n\n".join(pages), "pages": offsets, "metadata": {"source": name}}

        doc = load_bytes(
            text_pdf(REPORT, OUTLINE, title="Annual Report 2025"),
            "report.pdf",
            extractors={".pdf": batched(service, pages=1, at_once=1)},
            source="https://a.blob.core.windows.net/raw/report.pdf",
        )
        assert doc.metadata["pieces"] == 3
        assert [t for _o, t, _l in doc.headings] == [
            "Annual report 2025",
            "Message from the CEO",
            "Our strategy",
        ]
        assert doc.metadata["title"] == "Annual Report 2025"


class TestATitleAPersonWouldRecognise:
    @pytest.mark.parametrize(
        "given, kept",
        [
            ("2025 Annual Report", "2025 Annual Report"),
            ("Microsoft Word - Lending policy.docx", "Lending policy"),
            ("  Untitled ", None),
            ("PowerPoint Presentation", None),
            ("", None),
            (None, None),
        ],
    )
    def test_placeholders_and_printer_names_are_not_titles(self, given, kept):
        assert _useful_title(given) == kept


# ================================================================ PowerPoint, Excel, HTML, Markdown ===


class TestADecksNotesAndTitle:
    def deck(self, tmp_path, core_title=None):
        sys.path.insert(0, str(Path(__file__).parent))
        from test_table_and_deck_loaders import _slide, make_deck

        path = tmp_path / "deck.pptx"
        make_deck(
            path,
            {
                "slide1.xml": _slide("Bread", ["Sourdough is leavened by wild yeast."]),
                "slide2.xml": _slide("Rock", ["Basalt."]),
            },
        )
        rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide"
        body = '<p:sp><p:nvSpPr><p:cNvPr id="3" name="Notes"/><p:cNvSpPr/><p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr><p:txBody><a:p><a:r><a:t>Mention the rye loaf.</a:t></a:r></a:p><a:p><a:r><a:t>It takes twelve hours.</a:t></a:r></a:p></p:txBody></p:sp>'
        number = '<p:sp><p:nvSpPr><p:cNvPr id="4" name="Slide Number"/><p:cNvSpPr/><p:nvPr><p:ph type="sldNum" idx="5"/></p:nvPr></p:nvSpPr><p:txBody><a:p><a:r><a:t>1</a:t></a:r></a:p></p:txBody></p:sp>'
        with zipfile.ZipFile(path, "a") as z:
            z.writestr(
                "ppt/slides/_rels/slide1.xml.rels",
                f'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId9" Type="{rel}" Target="../notesSlides/notesSlide1.xml"/></Relationships>',
            )
            z.writestr(
                "ppt/notesSlides/notesSlide1.xml",
                f'<p:notes xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:cSld><p:spTree>{body}{number}</p:spTree></p:cSld></p:notes>',
            )
            if core_title:
                z.writestr(
                    "docProps/core.xml",
                    f'<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>{core_title}</dc:title></cp:coreProperties>',
                )
        return path

    def test_speaker_notes_follow_their_slide(self, tmp_path):
        text = load(self.deck(tmp_path)).text
        slide_one = text.split("# Slide 2")[0].rstrip()
        assert slide_one.endswith("Speaker notes: Mention the rye loaf.\nIt takes twelve hours."), (
            "and not the notes page's slide number"
        )
        assert "Speaker notes" not in text.split("# Slide 2")[1], "only the slide that has them"

    def test_the_decks_title_else_its_first_slides(self, tmp_path):
        assert (
            load(self.deck(tmp_path, core_title="Bakery basics")).metadata["title"]
            == "Bakery basics"
        )
        other = tmp_path / "other"
        other.mkdir()
        assert load(self.deck(other)).metadata["title"] == "Bread"


class TestAWorkbooksTitle:
    def test_the_workbooks_own_title(self, tmp_path):
        openpyxl = pytest.importorskip("openpyxl")
        wb = openpyxl.Workbook()
        wb.active.append(["Region", "Revenue"])
        wb.active.append(["EMEA", 1200])
        wb.properties.title = "Quarterly sales"
        wb.save(tmp_path / "book.xlsx")
        assert load(tmp_path / "book.xlsx").metadata["title"] == "Quarterly sales"


class TestAWebPage:
    PAGE = """<html><head><title>Offline guide</title></head><body>
<h1>Offline use</h1><p>Nothing downloads unless you ask.</p>
<table><caption>Speeds</caption><tr><th>Mode</th><th>Speed</th></tr><tr><td>dense</td><td>fast</td></tr>
<tr><td>hybrid</td><td><table><tr><td>fast</td><td>on a laptop</td></tr></table></td></tr></table>
<ul><li>No account</li><li>No network</li></ul><ol start="3"><li>Third</li><li>Fourth</li></ol>
</body></html>"""

    def test_tables_are_rows_lists_keep_their_markers_and_the_title_is_kept(self, tmp_path):
        (tmp_path / "page.html").write_text(self.PAGE, encoding="utf-8")
        doc = load(tmp_path / "page.html")
        assert "Speeds\nMode: dense; Speed: fast\nMode: hybrid; Speed: fast on a laptop" in doc.text
        assert "- No account\n- No network" in doc.text
        assert "3. Third\n4. Fourth" in doc.text
        assert doc.metadata["title"] == "Offline guide" and "Offline guide" not in doc.text

    def test_a_page_with_no_title_takes_its_first_heading(self, tmp_path):
        (tmp_path / "page.html").write_text(
            "<h1>Offline use</h1><p>Nothing downloads.</p>", encoding="utf-8"
        )
        assert load(tmp_path / "page.html").metadata["title"] == "Offline use"


class TestAMarkdownFilesTitle:
    def test_its_front_matter_else_its_first_top_level_heading(self, tmp_path):
        (tmp_path / "a.md").write_text(
            "---\ntitle: Lending terms\n---\n\n# Terms\n\nPay in thirty days.\n", encoding="utf-8"
        )
        (tmp_path / "b.md").write_text(
            "Intro.\n\n## Detail\n\n# Terms\n\nPay in thirty days.\n", encoding="utf-8"
        )
        (tmp_path / "c.md").write_text("Just words.\n", encoding="utf-8")
        assert load(tmp_path / "a.md").metadata["title"] == "Lending terms"
        assert load(tmp_path / "b.md").metadata["title"] == "Terms"
        assert "title" not in load(tmp_path / "c.md").metadata


# ================================================================ through the extraction app ===


class TestThroughTheExtractionApp:
    def test_a_word_files_table_and_title_reach_every_chunk(self, tmp_path, monkeypatch):
        docx = docx_module()
        from fastapi.testclient import TestClient

        from vectrixdb.api.extraction import ExtractionService, create_extraction_app
        from vectrixdb.extract import coerce

        monkeypatch.delenv("VECTRIXDB_API_KEY", raising=False)
        d = docx.Document()
        d.add_paragraph("Lending Policy 2025", style="Title")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Region", "Revenue"
        t.cell(1, 0).text, t.cell(1, 1).text = "EMEA", "1200"
        app = TestClient(create_extraction_app(ExtractionService(), allow_open=True))
        reply = app.post(
            "/extract/docx",
            content=saved(d, tmp_path).read_bytes(),
            headers={"X-Filename": "policy.docx", "Accept": "application/json"},
        ).json()
        assert (
            reply["metadata"]["title"] == "Lending Policy 2025"
            and "Region: EMEA; Revenue: 1200" in reply["text"]
        )
        prepared = prepare_document(
            coerce(reply, "policy.docx"),
            "policy.docx",
            chunk="markdown",
            chunk_size=1000,
            overlap=0,
        )
        assert all(m["title"] == "Lending Policy 2025" for m in prepared.metadata)
        assert any("Region: EMEA; Revenue: 1200" in text for text in prepared.texts)
