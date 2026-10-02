"""The structure of a document kept: numbers, headers, tables, charts, slides and Markdown, as the author wrote them.

The audit read files the way people make them and found the shape gone:
"1.2(b)" printed as "a.", a table of contents indexed as the first chunk,
"Heading Caption" taken for a heading, a Word table's header rows lost, a form
read as a header over one row, a slide deck read in the order its boxes were
added, a chart's numbers missing, "# install dependencies" in a bash block
taken for a heading, a title merged across columns read as a header, 0.042
where the sheet showed 4.2%, and a formula nobody recalculated read as nothing.
"""

from __future__ import annotations

import io
import zipfile

import pytest

from vectrixdb.ingest import (
    _chart_lines,
    _format_cell,
    _list_marker,
    _markdown_headings,
    _sheet_lines,
    load,
)


# ============================================================ tables anywhere ===


class TestWhereATableStartsAndWhatHeadsIt:
    def test_a_title_above_a_table_is_text_and_the_header_is_found_under_it(self):
        rows = [
            ["Quarterly Sales Report", None, None],
            [None, None, None],
            ["Region", "Year", "Revenue"],
            ["EMEA", 2024, 1200],
        ]
        assert _sheet_lines(rows) == [
            "Quarterly Sales Report",
            "Region: EMEA; Year: 2024; Revenue: 1200",
        ]

    def test_two_tables_on_one_sheet_each_keep_their_own_header(self):
        rows = [["Region", "Revenue"], ["EMEA", 1200], [], ["Employee", "Salary"], ["Ana", 70000]]
        assert _sheet_lines(rows) == ["Region: EMEA; Revenue: 1200", "Employee: Ana; Salary: 70000"]

    def test_tables_side_by_side_are_two_tables(self):
        rows = [["Region", "Revenue", None, "Product", "Units"], ["EMEA", 1200, None, "Laptop", 40]]
        assert _sheet_lines(rows) == ["Region: EMEA; Revenue: 1200", "Product: Laptop; Units: 40"]

    def test_a_header_over_two_rows_is_joined_per_column(self):
        rows = [
            ["Region", "2023", None, "2024", None],
            [None, "H1", "H2", "H1", "H2"],
            ["EMEA", 500, 600, 700, 800],
        ]
        assert _sheet_lines(rows) == [
            "Region: EMEA; 2023 H1: 500; 2023 H2: 600; 2024 H1: 700; 2024 H2: 800"
        ]

    def test_years_written_as_numbers_head_their_columns(self):
        assert _sheet_lines([["Region", 2023, 2024], ["EMEA", 1100, 1200]]) == [
            "Region: EMEA; 2023: 1100; 2024: 1200"
        ]

    def test_a_csv_with_no_header_is_not_given_one(self):
        rows = [["EMEA", "2024", "1200"], ["APAC", "2024", "950"]]
        assert _sheet_lines(rows) == ["EMEA; 2024; 1200", "APAC; 2024; 950"]

    def test_a_form_of_two_columns_is_keys_and_values(self):
        rows = [
            ["Applicant name", "Jane Doe"],
            ["Date of birth", "1980-04-02"],
            ["Account number", "12345678"],
        ]
        assert _sheet_lines(rows) == [
            "Applicant name: Jane Doe",
            "Date of birth: 1980-04-02",
            "Account number: 12345678",
        ]

    def test_a_table_that_goes_on_after_a_blank_line_keeps_its_header(self):
        rows = [["A", "B"], [], ["1", "2"], [None, None], [], ["3", "4"]]
        assert _sheet_lines(rows) == ["A: 1; B: 2", "A: 3; B: 4"]

    def test_a_title_merged_across_columns_is_one_line(self):
        rows = [["Quarterly Sales Report"] * 3, ["Region", "Year", "Revenue"], ["EMEA", 2024, 1200]]
        assert _sheet_lines(rows)[0] == "Quarterly Sales Report"

    def test_a_markdown_header_with_no_cell_over_the_labels_heads_the_figures(self):
        """A reading by sight wrote Table 15 this way, and 2025 headed the labels while 2024's figure had no year."""
        from vectrixdb.ingest import markdown_document

        text = "| 2025 | 2024 |\n|---|---|\n| Personal banking | $ 14,500 | $ 13,828 |\n| Total | $ 20,686 | $ 19,790 |"
        assert (
            markdown_document(text).text
            == "Personal banking; 2025: $ 14,500; 2024: $ 13,828\nTotal; 2025: $ 20,686; 2024: $ 19,790"
        )
        assert (
            markdown_document("| Region | Revenue |\n|---|---|\n| EMEA | 1200 |").text
            == "Region: EMEA; Revenue: 1200"
        )


class TestAValueAsTheSheetShowsIt:
    @pytest.mark.parametrize(
        "value,code,shown",
        [
            (0.042, "0.0%", "4.2%"),
            (1234.5, '"$"#,##0.00', "$1,234.50"),
            (-250, "#,##0.00;(#,##0.00)", "(250.00)"),
            (1200000, "#,##0", "1,200,000"),
            (99.9, "#,##0.00 [$€-x-euro2]", "99.90 €"),
            (0.0000123, "0.00E+00", "1.23E-05"),
            (42, "General", 42),
        ],
    )
    def test_formats(self, value, code, shown):
        assert _format_cell(value, code) == shown

    def test_a_month_is_a_month(self):
        import datetime

        assert _format_cell(datetime.datetime(2024, 3, 1), "mmm yyyy") == "Mar 2024"
        assert _format_cell(datetime.timedelta(hours=36), "[h]:mm") == "36:00"


# ============================================================ excel ===


def workbook(tmp_path, build, name="book.xlsx"):
    openpyxl = pytest.importorskip("openpyxl")
    book = openpyxl.Workbook()
    build(book)
    book.save(tmp_path / name)
    return tmp_path / name


class TestAWorkbookAsItReads:
    def test_merged_headers_formats_and_hidden_rows_and_columns(self, tmp_path):
        def build(book):
            sheet = book.active
            sheet.title = "Revenue"
            sheet.append(["Region", "2024", None, "Internal"])
            sheet.append([None, "H1", "H2", None])
            sheet.append(["EMEA", 0.042, 0.051, "X-1"])
            sheet.append(["APAC", 0.03, 0.02, "X-2"])
            sheet.merge_cells("B1:C1")
            for row in sheet.iter_rows(min_row=3, max_row=4, min_col=2, max_col=3):
                for cell in row:
                    cell.number_format = "0.0%"
            sheet.column_dimensions["D"].hidden = True
            sheet.row_dimensions[4].hidden = True

        text = load(workbook(tmp_path, build)).text
        assert "Region: EMEA; 2024 H1: 4.2%; 2024 H2: 5.1%" in text
        assert "X-1" not in text and "APAC" not in text

    def test_a_formula_with_no_value_saved_shows_its_formula(self, tmp_path):
        def build(book):
            sheet = book.active
            sheet.append(["Item", "Qty", "Price", "Total"])
            sheet.append(["Pen", 3, 1.5, "=B2*C2"])

        assert (
            "Item: Pen; Qty: 3; Price: 1.5; Total: =B2*C2" in load(workbook(tmp_path, build)).text
        )

    def test_a_charts_title_is_read(self, tmp_path):
        def build(book):
            from openpyxl.chart import BarChart, Reference

            sheet = book.active
            for row in (["Region", "Revenue"], ["EMEA", 41.7], ["APAC", 28.3]):
                sheet.append(row)
            chart = BarChart()
            chart.title = "Revenue by region"
            chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
            chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
            sheet.add_chart(chart, "D2")

        assert "Chart: Revenue by region (bar chart)" in load(workbook(tmp_path, build)).text


class TestAChartsNumbers:
    CHART = (
        '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:chart>'
        "<c:title><c:tx><c:rich><a:p><a:r><a:t>Share by region</a:t></a:r></a:p></c:rich></c:tx></c:title>"
        '<c:plotArea><c:pieChart><c:ser><c:tx><c:strRef><c:strCache><c:pt idx="0"><c:v>FY2024</c:v></c:pt></c:strCache></c:strRef></c:tx>'
        '<c:cat><c:strRef><c:strCache><c:pt idx="0"><c:v>EMEA</c:v></c:pt><c:pt idx="1"><c:v>APAC</c:v></c:pt></c:strCache></c:strRef></c:cat>'
        '<c:val><c:numRef><c:numCache><c:formatCode>0.0%</c:formatCode><c:pt idx="0"><c:v>0.417</c:v></c:pt>'
        '<c:pt idx="1"><c:v>0.283</c:v></c:pt></c:numCache></c:numRef></c:val></c:ser></c:pieChart></c:plotArea></c:chart></c:chartSpace>'
    ).encode()

    def test_a_chart_is_its_title_and_a_row_a_category(self):
        assert _chart_lines(self.CHART) == [
            "Chart: Share by region (pie chart)",
            "Category: EMEA; FY2024: 41.7%",
            "Category: APAC; FY2024: 28.3%",
        ]


# ============================================================ powerpoint ===


P_NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'


def shape(text_xml, x, y, placeholder=None):
    ph = f'<p:nvPr><p:ph type="{placeholder}"/></p:nvPr>' if placeholder else "<p:nvPr/>"
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="1" name="s"/><p:cNvSpPr/>{ph}</p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="100" cy="100"/></a:xfrm></p:spPr>'
        f"<p:txBody><a:bodyPr/>{text_xml}</p:txBody></p:sp>"
    )


def para(words, level=None, bullet=False):
    props = f'<a:pPr lvl="{level}">' if level is not None else "<a:pPr>"
    props += '<a:buChar char="&#8226;"/></a:pPr>' if bullet else "</a:pPr>"
    runs = "".join(f"<a:r><a:t>{w}</a:t></a:r>" if w != "\n" else "<a:br/>" for w in words)
    return f"<a:p>{props}{runs}</a:p>"


def deck(tmp_path, slides, name="deck.pptx"):
    path = tmp_path / name
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        ids = "".join(f'<p:sldId id="{256 + n}" r:id="rId{n}"/>' for n in range(1, len(slides) + 1))
        archive.writestr(
            "ppt/presentation.xml",
            f"<p:presentation {P_NS}><p:sldIdLst>{ids}</p:sldIdLst></p:presentation>",
        )
        rels = "".join(
            f'<Relationship Id="rId{n}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide{n}.xml"/>'
            for n in range(1, len(slides) + 1)
        )
        archive.writestr(
            "ppt/_rels/presentation.xml.rels",
            f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rels}</Relationships>',
        )
        for n, (body, show) in enumerate(slides, start=1):
            hidden = ' show="0"' if not show else ""
            archive.writestr(
                f"ppt/slides/slide{n}.xml",
                f"<p:sld {P_NS}{hidden}><p:cSld><p:spTree>{body}</p:spTree></p:cSld></p:sld>",
            )
    return path


class TestASlideAsItReads:
    def test_boxes_are_read_by_where_they_sit_not_the_order_they_were_added(self, tmp_path):
        body = (
            shape(para(["4. Conclusion"]), 0, 4000000)
            + shape(para(["1. Introduction"]), 0, 1000000)
            + shape(para(["Next steps"]), 0, 100000, "title")
            + shape(para(["3. Risks"]), 0, 3000000)
            + shape(para(["2. Options"]), 0, 2000000)
        )
        text = load(deck(tmp_path, [(body, True)])).text
        assert (
            text.index("1. Introduction")
            < text.index("2. Options")
            < text.index("3. Risks")
            < text.index("4. Conclusion")
        )
        assert text.startswith("# Slide 1: Next steps")

    def test_bullet_levels_line_breaks_and_no_footer_date_or_number(self, tmp_path):
        body = (
            shape(para(["Strategy"]), 0, 0, "title")
            + shape(
                para(["Grow deposits"], 0, True)
                + para(["Retail"], 1, True)
                + para(["Close 4 branches", "\n", "Merge back office"], 0, True),
                0,
                1000000,
                "body",
            )
            + shape(para(["Confidential - internal"]), 0, 6000000, "ftr")
            + shape(para(["3/4/2025"]), 0, 6000000, "dt")
            + shape(para(["7"]), 0, 6000000, "sldNum")
        )
        text = load(deck(tmp_path, [(body, True)])).text
        assert "- Grow deposits\n  - Retail\n- Close 4 branches\n- Merge back office" in text
        assert "Confidential" not in text and "3/4/2025" not in text

    def test_a_hidden_slide_is_named_and_not_read(self, tmp_path):
        doc = load(
            deck(
                tmp_path,
                [
                    (shape(para(["Shown"]), 0, 0, "title"), True),
                    (shape(para(["Draft only"]), 0, 0, "title"), False),
                ],
            )
        )
        assert "Draft only" not in doc.text and doc.metadata["slides_hidden"] == [2]

    def test_a_pictures_alt_text_and_a_title_in_a_plain_box(self, tmp_path):
        picture = (
            '<p:pic><p:nvPicPr><p:cNvPr id="4" name="Picture 3" descr="Map of 42 branch locations across Ontario"/>'
            '<p:cNvPicPr/><p:nvPr/></p:nvPicPr><p:spPr><a:xfrm><a:off x="0" y="2000000"/></a:xfrm></p:spPr></p:pic>'
        )
        text = load(deck(tmp_path, [(shape(para(["Branch network"]), 0, 0) + picture, True)])).text
        assert (
            text.startswith("# Slide 1: Branch network")
            and "[Figure: Map of 42 branch locations across Ontario]" in text
        )


# ============================================================ word ===


def document_bytes(build):
    docx = pytest.importorskip("docx")
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def with_body(data, body_xml):
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "word/document.xml":
                text = content.decode("utf-8")
                start, end = text.index("<w:body>") + len("<w:body>"), text.index("<w:sectPr")
                content = (text[:start] + body_xml + text[end:]).encode("utf-8")
            target.writestr(item, content)
    return out.getvalue()


class TestWordAsWordShowsIt:
    def test_list_numbers_as_printed(self):
        lists = {
            "7": {
                0: ("decimal", 1, "%1.", False),
                1: ("decimal", 1, "%1.%2", False),
                2: ("lowerLetter", 1, "(%3)", False),
            }
        }
        counters: dict = {}
        shown = [
            _list_marker(("7", level), lists, counters).strip() for level in (0, 1, 1, 2, 2, 0, 1)
        ]
        assert shown == ["1.", "1.1", "1.2", "(a)", "(b)", "2.", "2.1"]
        roman = {"9": {0: ("upperRoman", 1, "Article %1 -", False)}}
        assert _list_marker(("9", 0), roman, {}).strip() == "Article I -"

    def test_a_numbered_list_keeps_its_numbers(self, tmp_path):
        def build(document):
            document.add_paragraph("First step", style="List Number")
            document.add_paragraph("Second step", style="List Number")

        (tmp_path / "l.docx").write_bytes(document_bytes(build))
        assert "1. First step\n\n2. Second step" in load(tmp_path / "l.docx").text

    def test_a_table_of_contents_is_not_read(self, tmp_path):
        def build(document):
            from docx.enum.style import WD_STYLE_TYPE

            document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
            document.add_paragraph("Purpose\t1", style="TOC 1")
            document.add_heading("Purpose", level=1)
            document.add_paragraph("The policy sets the buffer.")

        (tmp_path / "t.docx").write_bytes(document_bytes(build))
        text = load(tmp_path / "t.docx").text
        assert text.startswith("# Purpose") and "\t1" not in text

    def test_a_style_that_only_starts_with_heading_is_not_a_heading(self, tmp_path):
        def build(document):
            from docx.enum.style import WD_STYLE_TYPE

            document.styles.add_style("Heading Caption", WD_STYLE_TYPE.PARAGRAPH)
            document.add_heading("Results", level=1)
            document.add_paragraph("Source: HR survey 2025", style="Heading Caption")
            document.add_heading("Details", level=2)

        (tmp_path / "h.docx").write_bytes(document_bytes(build))
        assert [h[1] for h in load(tmp_path / "h.docx").headings] == ["Results", "Details"]

    def test_the_rows_word_marks_as_the_header_head_the_table(self, tmp_path):
        row = lambda cells, head=False: (  # noqa: E731
            "<w:tr>"
            + ("<w:trPr><w:tblHeader/></w:trPr>" if head else "")
            + "".join(f"<w:tc><w:p><w:r><w:t>{c}</w:t></w:r></w:p></w:tc>" for c in cells)
            + "</w:tr>"
        )
        body = (
            "<w:tbl>"
            + row(["Region", "Revenue"], True)
            + row(["", "2025"], True)
            + row(["EMEA", "100"])
            + "</w:tbl>"
        )
        (tmp_path / "r.docx").write_bytes(with_body(document_bytes(lambda d: None), body))
        assert "Region: EMEA; Revenue 2025: 100" in load(tmp_path / "r.docx").text

    def test_a_form_table_is_keys_and_values(self, tmp_path):
        def build(document):
            table = document.add_table(rows=3, cols=2)
            for row, (key, value) in zip(
                table.rows,
                [
                    ("Applicant name", "Jane Doe"),
                    ("Date of birth", "1980-04-02"),
                    ("Account number", "12345678"),
                ],
            ):
                row.cells[0].text, row.cells[1].text = key, value

        (tmp_path / "f.docx").write_bytes(document_bytes(build))
        assert (
            "Applicant name: Jane Doe\nDate of birth: 1980-04-02" in load(tmp_path / "f.docx").text
        )

    def test_headers_and_footers_are_the_documents_and_not_its_text(self, tmp_path):
        def build(document):
            section = document.sections[0]
            section.header.paragraphs[0].text = "Board Pack Q3 2025: Liquidity Review"
            section.footer.paragraphs[0].text = "CONFIDENTIAL"
            document.add_paragraph("Liquidity remained strong through the quarter.")

        doc = (
            load(tmp_path / "hf.docx")
            if (tmp_path / "hf.docx").write_bytes(document_bytes(build))
            else None
        )
        assert doc.metadata["page_header"] == "Board Pack Q3 2025: Liquidity Review"
        assert doc.metadata["page_footer"] == "CONFIDENTIAL"
        assert doc.metadata["title"] == "Board Pack Q3 2025: Liquidity Review", (
            "no title of its own and no heading: its header"
        )
        assert "CONFIDENTIAL" not in doc.text

    def test_an_equation_in_a_line_and_a_comment_kept(self, tmp_path):
        math = (
            '<w:p><m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            "<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b+c</m:t></m:r></m:den></m:f>"
            "<m:r><m:t>=</m:t></m:r><m:sSup><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>"
            "</m:oMath></w:p>"
        )
        (tmp_path / "e.docx").write_bytes(with_body(document_bytes(lambda d: None), math))
        assert "a/(b+c)=x^2" in load(tmp_path / "e.docx").text

        def build(document):
            paragraph = document.add_paragraph("Invoices are paid within thirty days.")
            document.add_comment(paragraph.runs, text="We cannot accept thirty days.", author="CFO")

        (tmp_path / "c.docx").write_bytes(document_bytes(build))
        assert "[Comment: CFO: We cannot accept thirty days.]" in load(tmp_path / "c.docx").text


# ============================================================ markdown ===


class TestMarkdownAsCommonMarkReadsIt:
    def test_a_comment_in_code_is_not_a_heading(self):
        text = "# Install\n\n```bash\n# install dependencies first\npip install x\n```\n\n## Configure\n"
        assert [h[1] for h in _markdown_headings(text)] == ["Install", "Configure"]

    def test_closing_hashes_need_a_space_and_a_bare_hash_is_nothing(self):
        assert [
            h[1]
            for h in _markdown_headings(
                "# Why C#\n\n## F# and C# compared ##\n\n#\nOrphan line\n\n#hashtag\n"
            )
        ] == ["Why C#", "F# and C# compared"]

    @pytest.mark.parametrize(
        "raw,title,headings",
        [
            (
                "Annual Report\n=============\n\nIntro.\n\nRevenue\n-------\n\nGrew.\n",
                "Annual Report",
                ["Annual Report", "Revenue"],
            ),
            (
                "<h1>From HTML</h1>\n\nText.\n\n<h2>Part <b>two</b></h2>\n\nMore.\n",
                "From HTML",
                ["From HTML", "Part two"],
            ),
            (
                '+++\ntitle = "TOML Front"\n+++\n\n# Real Heading\n\nBody.\n',
                "TOML Front",
                ["Real Heading"],
            ),
        ],
        ids=["setext", "html headings", "toml front matter"],
    )
    def test_other_ways_of_writing_a_heading_or_a_title(self, tmp_path, raw, title, headings):
        (tmp_path / "d.md").write_text(raw, encoding="utf-8")
        doc = load(tmp_path / "d.md")
        assert doc.metadata["title"] == title and [h[1] for h in doc.headings] == headings

    def test_images_in_links_with_brackets_and_by_reference(self, tmp_path):
        raw = (
            "# Links\n\n[![build badge](https://ci.example.test/badge.svg)](https://ci.example.test)\n\n"
            "![Chart](img/chart_(final).png)\n\n![Ref image][imgref]\n\n[imgref]: img/ref.png\n\nLine with<br>break and <b>bold</b>.\n"
        )
        (tmp_path / "l.md").write_text(raw, encoding="utf-8")
        doc = load(tmp_path / "l.md")
        assert "[Figure: build badge]" in doc.text and "](https://ci.example.test)" not in doc.text
        assert "[Figure: Chart]" in doc.text and "final" not in doc.text.replace(
            "[Figure: Chart]", ""
        )
        assert "[Figure: Ref image]" in doc.text and "[imgref]:" not in doc.text
        assert "Line with\nbreak and bold." in doc.text
