"""Pictures inside Word, PowerPoint and Excel files, found where they sit and described like a PDF's.

Before, those three readers took the text and nothing else, so a chart pasted
into a report was not in the index at all. With ``images=True`` each picture
becomes a figure line where it sits, which is what a describer writes under.
Without it the text is exactly what it always was.
"""

from __future__ import annotations

import io
import random
import zipfile

import pytest

Image = pytest.importorskip("PIL.Image")

from vectrixdb.ingest import describe_figures, load  # noqa: E402

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"


def picture(seed: int = 0, size=(120, 90)) -> bytes:
    """A real PNG with something in it: big enough, and busy enough, not to be taken for decoration."""
    rng = random.Random(seed)
    image = Image.frombytes("RGB", size, bytes(rng.randrange(256) for _ in range(size[0] * size[1] * 3)))
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


def word(tmp_path, *pictures: bytes):
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_heading("Results", 1)
    document.add_paragraph("Revenue grew in every region.")
    for data in pictures:
        document.add_picture(io.BytesIO(data))
    document.add_paragraph("The chart above is unaudited.")
    path = tmp_path / "report.docx"
    document.save(path)
    return path


def deck(tmp_path, slides, external: bool = False):
    """A deck written by hand, one slide per entry of (title, line, pictures), so no package is needed."""
    path = tmp_path / "deck.pptx"
    with zipfile.ZipFile(path, "w") as archive:
        ids = "".join(f'<p:sldId id="{255 + n}" r:id="rId{n}"/>' for n in range(1, len(slides) + 1))
        archive.writestr("ppt/presentation.xml", f'<p:presentation xmlns:p="{P}" xmlns:r="{R}"><p:sldIdLst>{ids}</p:sldIdLst></p:presentation>')
        links = "".join(f'<Relationship Id="rId{n}" Type="slide" Target="slides/slide{n}.xml"/>' for n in range(1, len(slides) + 1))
        archive.writestr("ppt/_rels/presentation.xml.rels", f'<Relationships xmlns="{REL}">{links}</Relationships>')
        media = 0
        for number, (title, line, shown) in enumerate(slides, start=1):
            pics, rels = "", ""
            for index, data in enumerate(shown, start=1):
                media += 1
                pics += f'<p:pic><p:blipFill><a:blip r:embed="rIdP{index}"/></p:blipFill></p:pic>'
                rels += f'<Relationship Id="rIdP{index}" Type="image" Target="../media/image{media}.png"/>'
                archive.writestr(f"ppt/media/image{media}.png", data)
            if external:
                pics += '<p:pic><p:blipFill><a:blip r:embed="rIdWeb"/></p:blipFill></p:pic>'
                rels += '<Relationship Id="rIdWeb" Type="image" Target="https://example.com/logo.png" TargetMode="External"/>'
            archive.writestr(
                f"ppt/slides/slide{number}.xml",
                f'<p:sld xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}"><p:cSld><p:spTree>'
                f'<p:sp><p:nvSpPr><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr><p:txBody><a:p><a:r><a:t>{title}</a:t></a:r></a:p></p:txBody></p:sp>'
                f"<p:sp><p:txBody><a:p><a:r><a:t>{line}</a:t></a:r></a:p></p:txBody></p:sp>{pics}"
                "</p:spTree></p:cSld></p:sld>",
            )
            archive.writestr(f"ppt/slides/_rels/slide{number}.xml.rels", f'<Relationships xmlns="{REL}">{rels}</Relationships>')
    return path


def workbook(tmp_path, with_picture_only_sheet: bool = True):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.drawing.image import Image as Pasted

    book = openpyxl.Workbook()
    first = book.active
    first.title = "Summary"
    first.append(["Region", "Revenue"])
    first.append(["EMEA", 1200])
    chart = book.create_sheet("Chart")
    chart.append(["Quarter", "Sales"])
    chart.append(["Q1", 10])
    chart.add_image(Pasted(io.BytesIO(picture(1))), "D2")
    if with_picture_only_sheet:
        book.create_sheet("Picture only").add_image(Pasted(io.BytesIO(picture(2))), "A1")
    path = tmp_path / "book.xlsx"
    book.save(path)
    return path


def described(data, context):
    return "A bar chart of revenue by region, EMEA highest."


class TestWord:
    def test_a_picture_is_a_figure_after_the_paragraph_that_shows_it(self, tmp_path):
        doc = load(word(tmp_path, picture()), images=True)
        assert doc.text == "# Results\n\nRevenue grew in every region.\n\n[Figure: fig1.png]\n\nThe chart above is unaudited."
        assert doc.figures == [(doc.text.index("[Figure"), {"caption": "fig1.png", "src": "fig1.png", "described": False})]
        assert doc.images["fig1.png"][:8] == b"\x89PNG\r\n\x1a\n"

    def test_the_pictures_are_numbered_in_the_order_they_are_shown(self, tmp_path):
        doc = load(word(tmp_path, picture(1), picture(2)), images=True)
        assert [info["src"] for _o, info in doc.figures] == ["fig1.png", "fig2.png"]

    def test_without_images_the_text_is_what_it_always_was(self, tmp_path):
        doc = load(word(tmp_path, picture()))
        assert doc.text == "# Results\n\nRevenue grew in every region.\n\nThe chart above is unaudited."
        assert doc.figures == [] and doc.images == {}

    def test_the_headings_still_point_at_their_words(self, tmp_path):
        doc = load(word(tmp_path, picture()), images=True)
        assert [(doc.text[o : o + len(h) + 2], h) for o, h, _level in doc.headings] == [("# Results", "Results")]


class TestPowerPoint:
    def test_a_picture_is_a_figure_at_the_end_of_its_slide(self, tmp_path):
        doc = load(deck(tmp_path, [("Revenue", "Up everywhere.", [picture()]), ("Costs", "Flat.", [])]), images=True)
        assert doc.text.startswith("# Slide 1: Revenue\n\nUp everywhere.\n\n[Figure: s1-fig1.png]\n\n# Slide 2: Costs")
        assert [info["src"] for _o, info in doc.figures] == ["s1-fig1.png"]

    def test_the_slides_are_still_the_pages(self, tmp_path):
        doc = load(deck(tmp_path, [("Revenue", "Up.", [picture()]), ("Costs", "Flat.", [])]), images=True)
        assert [doc.text[o : o + 9] for o, _n in doc.pages] == ["# Slide 1", "# Slide 2"]

    def test_a_picture_linked_from_the_web_is_not_fetched(self, tmp_path):
        doc = load(deck(tmp_path, [("Revenue", "Up.", [])], external=True), images=True)
        assert doc.figures == [] and doc.images == {}

    def test_without_images_the_text_is_what_it_always_was(self, tmp_path):
        assert "[Figure" not in load(deck(tmp_path, [("Revenue", "Up.", [picture()])])).text


class TestExcel:
    def test_a_picture_is_a_figure_after_its_sheets_rows(self, tmp_path):
        doc = load(workbook(tmp_path), images=True)
        chart = doc.text.index("# Chart")
        assert doc.text.index("Quarter: Q1; Sales: 10") < doc.text.index("[Figure: sheet2-fig1.png]") > chart

    def test_a_sheet_that_is_only_a_picture_is_still_read(self, tmp_path):
        doc = load(workbook(tmp_path), images=True)
        assert "# Picture only\n\n[Figure: sheet3-fig1.png]" in doc.text
        assert doc.metadata["sheets"] == 3

    def test_without_images_the_text_is_what_it_always_was(self, tmp_path):
        doc = load(workbook(tmp_path))
        assert "[Figure" not in doc.text and "Picture only" not in doc.text and doc.metadata["sheets"] == 2


class TestTheyAreDescribedLikeAPdfsFigures:
    def test_the_description_goes_under_the_figure_where_it_sits(self, tmp_path):
        doc = describe_figures(load(word(tmp_path, picture()), images=True), described, name="report.docx")
        figure = doc.text.index("[Figure:")
        assert doc.text.index("A bar chart of revenue by region") > figure > doc.text.index("Revenue grew")
        assert doc.figures[0][1]["described"] is True

    def test_a_picture_on_every_slide_is_a_letterhead_and_is_never_described(self, tmp_path):
        """The same picture three times is decoration, and it is not paid for."""
        logo = picture(7)
        asked = []
        doc = load(deck(tmp_path, [(f"Slide {n}", "Text.", [logo]) for n in range(1, 4)]), images=True)
        doc = describe_figures(doc, lambda data, context: asked.append(context) or "a logo")
        assert asked == [] and "[Figure" not in doc.text
